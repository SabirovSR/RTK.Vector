import argparse
import getpass
import os
from pathlib import Path
from datetime import timedelta
from sqlalchemy import select
from .db import SessionLocal, now
from .models import *
from .security import hasher


def seed(db):
    if db.scalar(select(User).where(User.email == "school@demo.test")):
        print("Демоданные уже загружены")
        return
    password = os.getenv("DEMO_PASSWORD", "VectorDemo2026!")
    school = User(
        email="school@demo.test",
        name="Анна Соколова",
        role="school",
        password=hasher.hash(password),
    )
    colleague = User(
        email="colleague@demo.test",
        name="Максим Волков",
        role="school",
        password=hasher.hash(password),
    )
    db.add_all([school, colleague])
    db.flush()
    unis = []
    for i, (name, region, profile) in enumerate(
        [
            ("Северный технологический университет", "Москва", "Технический"),
            ("Волжский университет цифровых технологий", "Татарстан", "ИТ и экономика"),
            ("Уральский инженерный институт", "Свердловская область", "Инженерный"),
            ("Балтийский университет", "Санкт-Петербург", "Многопрофильный"),
        ]
    ):
        u = University(
            name=name,
            region=region,
            profile=profile,
            accreditation="Действует",
            owner_id=school.id if i < 3 else colleague.id,
            details={
                "departments": "Институт информационных технологий",
                "directions": "Программирование, анализ данных",
                "requisites": "Синтетическая организация для демонстрации",
            },
        )
        db.add(u)
        db.flush()
        unis.append(u)
        db.add(
            Contact(
                university_id=u.id,
                name=["Ирина Лебедева", "Олег Миронов", "Вера Орлова", "Павел Сергеев"][
                    i
                ],
                email=f"contact{i}@example.test",
                phone=f"+7000000000{i}",
                position="Руководитель ИТ-направления",
            )
        )
    db.add_all(
        [
            User(
                email="university@demo.test",
                name="Ирина Лебедева",
                role="university",
                university_id=unis[0].id,
                password=hasher.hash(password),
            ),
            User(
                email="other@demo.test",
                name="Олег Миронов",
                role="university",
                university_id=unis[1].id,
                password=hasher.hash(password),
            ),
        ]
    )
    names = [
        "Анализ данных без программирования",
        "Инженер-тестировщик",
        "Управление ИТ-проектами",
        "Промпт-инжиниринг",
        "Python-разработчик с инструментами ИИ",
    ]
    programs = []
    for i, name in enumerate(names):
        p = Program(
            name=name,
            direction=[
                "Аналитика данных",
                "Тестирование",
                "Управление",
                "Искусственный интеллект",
                "Разработка",
            ][i],
            competencies="Практические навыки, командный проект, итоговая оценка",
            tools=["RT.DataVision", "Яга", "Яга", "Web3Gate", "Python"][i],
        )
        db.add(p)
        db.flush()
        programs.append(p)
    stages = [
        "approval",
        "preparation",
        "training",
        "qualification",
        "new",
        "completed",
    ]
    for i, stage in enumerate(stages):
        uni = unis[0] if i in (0, 1, 2, 5) else unis[i % 3]
        deal = Deal(
            university_id=uni.id,
            owner_id=school.id,
            program_id=programs[i % 5].id,
            title=names[i % 5] + " · осенний набор",
            stage=stage,
            deadline=(now() + timedelta(days=14 + i * 3)).date().isoformat(),
            qualification={
                "contact": "Руководитель направления",
                "interest": "Набор практических групп",
                "budget": "Согласован",
                "window": "Осенний семестр",
            },
            preparation={
                k: "done" for k in ("materials", "licenses", "teachers", "curriculum")
            },
            notes="Внутренняя заметка ИТ-школы. Синтетические данные.",
            created_at=now() - timedelta(days=10 + i * 4),
            closed_at=now() if stage == "completed" else None,
        )
        db.add(deal)
        db.flush()
        db.add(
            Proposal(
                deal_id=deal.id,
                version=1,
                content=f"{names[i % 5]}\n72 академических часа. Практика на отечественных ИТ-продуктах.\nКомпетенции: анализ задач, работа с инструментами, командный проект.\nИтог: оценка компетенций и сертификат.",
                status="pending"
                if stage in ("approval", "new", "qualification")
                else "approved",
            )
        )
        db.add(
            Activity(
                deal_id=deal.id,
                actor_id=school.id,
                kind="comment",
                text="Подготовили программу для вашего вуза. Будем рады обсудить состав группы и сроки.",
                shared=True,
            )
        )
        db.add(
            Task(
                deal_id=deal.id,
                owner_id=school.id,
                title=[
                    "Получить решение по программе",
                    "Проверить готовность к запуску",
                    "Обсудить прогресс обучения",
                ][i % 3],
                due=(now() + timedelta(days=i - 1)).date().isoformat(),
            )
        )
        if stage in ("preparation", "training", "completed"):
            group = Group(
                deal_id=deal.id,
                name=f"Поток {i + 1} · 2026",
                confirmed=stage != "preparation",
                started=stage in ("training", "completed"),
            )
            db.add(group)
            db.flush()
            for j, name in enumerate(
                ["Александра Демо", "Даниил Примеров", "Софья Учебная"]
            ):
                person = Participant(
                    group_id=group.id,
                    name=name,
                    email=f"student{i}-{j}@example.test",
                    phone="+70000000000",
                )
                db.add(person)
                db.flush()
                if group.started:
                    db.add(
                        Result(
                            participant_id=person.id,
                            progress=100 if stage == "completed" else 40 + j * 15,
                            attendance=90 + j,
                            score=80 + j,
                        )
                    )
            if group.started:
                from .main import STORAGE

                STORAGE.mkdir(parents=True, exist_ok=True)
                from reportlab.pdfgen.canvas import Canvas

                path = f"demo-contract-{deal.id}.pdf"
                canvas = Canvas(str(STORAGE / path))
                canvas.drawString(50, 750, "DEMO signed agreement - synthetic document")
                canvas.save()
                db.add(
                    Document(
                        deal_id=deal.id,
                        name="demo-contract.pdf",
                        path=path,
                        kind="signed_contract",
                        version=1,
                        uploader_id=school.id,
                    )
                )
    other_deal = Deal(
        university_id=unis[3].id,
        owner_id=colleague.id,
        program_id=programs[0].id,
        title="Аналитика · новый партнёр",
    )
    db.add(other_deal)
    db.flush()
    db.add(
        Task(
            deal_id=other_deal.id,
            owner_id=colleague.id,
            title="Первая встреча с вузом",
            due=now().date().isoformat(),
        )
    )
    db.add(
        Notification(
            user_id=school.id,
            text="Добро пожаловать! Это синтетические демонстрационные данные.",
        )
    )
    print("Демоданные загружены. Логины: school@demo.test / university@demo.test")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["seed", "create-manager", "import-vendors"])
    parser.add_argument("--email")
    parser.add_argument("--name")
    parser.add_argument("--file")
    parser.add_argument("--password-stdin", action="store_true")
    args = parser.parse_args()
    with SessionLocal() as db:
        if args.command == "seed":
            seed(db)
        elif args.command == "create-manager":
            if not args.email or not args.name:
                parser.error("Нужны --email и --name")
            if args.password_stdin:
                import sys

                password = sys.stdin.readline().rstrip("\r\n")
            else:
                password = getpass.getpass("Пароль (не менее 10 символов): ")
            if len(password) < 10:
                parser.error("Пароль слишком короткий")
            db.add(
                User(
                    email=args.email.lower(),
                    name=args.name,
                    role="school",
                    password=hasher.hash(password),
                )
            )
        else:
            from openpyxl import load_workbook

            if not args.file:
                parser.error("Нужен --file")
            book = load_workbook(args.file, read_only=True, data_only=True)
            for company, product, name, phone, email, channel, *_ in list(
                book.active.values
            )[1:]:
                if not company or not product:
                    continue
                if not db.scalar(
                    select(Vendor).where(
                        Vendor.company == company, Vendor.product == product
                    )
                ):
                    db.add(
                        Vendor(
                            company=company,
                            product=product,
                            contact={
                                "name": name,
                                "phone": phone,
                                "email": email,
                                "channel": channel,
                            },
                        )
                    )
            book.close()
        db.commit()


if __name__ == "__main__":
    main()
