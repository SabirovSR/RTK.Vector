import io
import json
import re
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy import select
from .models import Participant, Application
from .schemas import ParticipantInput
from .security import fail


def phone(value):
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    return "+" + digits if digits else ""


def parse_import(data, filename):
    rows, errors, ignored = [], [], []
    try:
        if filename.lower().endswith(".json"):
            source = json.loads(data.decode("utf-8-sig"))
            if not isinstance(source, list):
                fail("Ожидается JSON-массив")
            if len(source) > 5000:
                fail("Не более 5000 записей за импорт")
            for index, value in enumerate(source, 1):
                if value is None:
                    continue
                if not isinstance(value, dict):
                    errors.append({"row": index, "message": "Ожидается объект заявки"})
                    continue
                rows.append(
                    {
                        "row": index,
                        "name": " ".join(
                            str(value.get(k) or "")
                            for k in ("Фамилия", "Имя", "Отчество")
                        ).strip(),
                        "email": str(value.get("Email") or "").lower().strip(),
                        "phone": phone(value.get("Телефон")),
                        "external_id": value.get("Номер заявки"),
                        "course": value.get("Курс", ""),
                    }
                )
                if not value.get("Номер заявки"):
                    errors.append({"row": index, "message": "Нет номера заявки"})
                if not isinstance(value.get("Номер заявки"), str) or not isinstance(
                    value.get("Курс"), str
                ):
                    errors.append(
                        {
                            "row": index,
                            "message": "Номер заявки и курс должны быть строками",
                        }
                    )
        elif filename.lower().endswith(".xlsx"):
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = book.worksheets[0]
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(x or "").strip() for x in next(iterator)]
            required = ["Фамилия", "Имя", "Email"]
            if any(k not in headers for k in required):
                fail("Нужны столбцы Фамилия, Имя и Email")
            allowed = {
                "Фамилия",
                "Имя",
                "Отчество",
                "Отчествопри наличии)",
                "Номер телефона",
                "Email",
            }
            ignored = [h for h in headers if h and h not in allowed]
            for index, row in enumerate(iterator, 2):
                if index > 5001:
                    fail("Не более 5000 строк за импорт")
                if not any(x is not None for x in row):
                    continue
                v = dict(zip(headers, row))
                rows.append(
                    {
                        "row": index,
                        "name": " ".join(
                            str(v.get(k) or "") for k in ("Фамилия", "Имя")
                        ).strip()
                        + " "
                        + str(v.get("Отчество") or v.get("Отчествопри наличии)") or ""),
                        "email": str(v.get("Email") or "").strip().lower(),
                        "phone": phone(v.get("Номер телефона")),
                    }
                )
            book.close()
        else:
            fail("Импорт поддерживает XLSX и JSON")
    except (
        ValueError,
        KeyError,
        StopIteration,
        OSError,
        BadZipFile,
        ParseError,
        InvalidFileException,
    ):
        fail("Не удалось прочитать файл импорта")
    if len(rows) > 5000:
        fail("Не более 5000 записей за импорт")
    for row in rows:
        try:
            item = ParticipantInput(
                name=row["name"].strip(), email=row["email"], phone=row["phone"]
            )
            row.update(item.model_dump())
        except ValidationError:
            errors.append({"row": row["row"], "message": "Проверьте ФИО и email"})
    return rows, errors, ignored


def prepare_import(db, group, rows, errors):
    seen, orders = set(), set()
    for row in rows:
        email = row["email"]
        existing = db.scalar(
            select(Participant).where(
                Participant.group_id == group.id, Participant.email == email
            )
        )
        row["action"] = "skip" if existing else "create"
        if email in seen:
            errors.append({"row": row["row"], "message": "Повтор email в файле"})
        seen.add(email)
        if existing and (
            existing.name.casefold() != row["name"].casefold()
            or (existing.phone and row["phone"] and existing.phone != row["phone"])
        ):
            errors.append(
                {
                    "row": row["row"],
                    "message": "Неоднозначный дубль: email совпадает, ФИО или телефон отличаются",
                }
            )
        external = row.get("external_id")
        if external is not None and not isinstance(external, str):
            continue
        if external:
            application = db.scalar(
                select(Application).where(Application.external_id == external)
            )
            if external in orders:
                errors.append({"row": row["row"], "message": "Повтор номера заявки"})
            orders.add(external)
            if application and (
                application.group_id != group.id
                or not existing
                or application.participant_id != existing.id
                or application.course != row["course"]
            ):
                errors.append(
                    {
                        "row": row["row"],
                        "message": "Номер заявки уже связан с другой записью",
                    }
                )


def apply_import(db, group, rows):
    count = 0
    for row in rows:
        person = db.scalar(
            select(Participant).where(
                Participant.group_id == group.id, Participant.email == row["email"]
            )
        )
        if not person:
            person = Participant(
                group_id=group.id,
                name=row["name"],
                email=row["email"],
                phone=row["phone"],
                kind=row.get("kind", "student"),
            )
            db.add(person)
            db.flush()
            count += 1
        if row.get("external_id") and not db.scalar(
            select(Application).where(Application.external_id == row["external_id"])
        ):
            db.add(
                Application(
                    external_id=row["external_id"],
                    group_id=group.id,
                    participant_id=person.id,
                    course=row["course"],
                )
            )
    if count:
        group.confirmed = False
    return count
